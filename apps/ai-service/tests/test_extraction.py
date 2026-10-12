import base64
import io

import pytest

from extraction import readers
from extraction.essays import (
    SheetMapping,
    apply_mapping,
    essay_from_images,
    guess_mapping,
    map_sheet,
    normalize_task_type,
    split_document,
)
from extraction.llm import PageTranscript, StubExtractor
from extraction.worker import ExtractionWorker, extract

ESSAY = " ".join(["Public transport reduces traffic and pollution in large cities."] * 6)


class FakeLLM:
    """Replays scripted JSON answers and page transcripts."""

    name = "fake"

    def __init__(self, answers=(), pages=None):
        self.answers = list(answers)
        self.pages = pages
        self.json_calls = []
        self.transcribe_calls = 0

    async def complete_json(self, system, user, schema):
        self.json_calls.append(user)
        answer = self.answers.pop(0)
        if isinstance(answer, Exception):
            raise answer
        return answer

    async def transcribe(self, images):
        self.transcribe_calls += 1
        return PageTranscript(pages=self.pages or [f"page {i}" for i in range(len(images))], task_prompt="Q?")


# ------------------------------------------------------------------ readers


def test_read_docx_returns_paragraph_text():
    import docx

    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    document.save(buf)

    assert readers.read_docx(buf.getvalue()) == "First paragraph.\nSecond paragraph."


def test_read_pdf_keeps_text_pages_and_renders_scans():
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "This page has a real text layer with enough characters.")
    doc.new_page()  # blank: treated as a scan
    pages = readers.read_pdf(doc.tobytes())

    assert "real text layer" in pages[0].text and pages[0].image is None
    assert pages[1].text == "" and pages[1].image.startswith(b"\x89PNG")


def test_read_sheets_csv():
    data = "Name,Essay\nAn,Hello there\n".encode()
    assert readers.read_sheets(data, "csv") == {"Sheet1": [["Name", "Essay"], ["An", "Hello there"]]}


# --------------------------------------------------------------- spreadsheets

SHEETS = {
    "Info": [["Class 10A"], ["Teacher: Ms Lan"]],
    "Bài làm": [
        ["DANH SÁCH BÀI VIẾT", "", ""],
        ["Họ tên", "Đề", "Bài viết"],
        ["Nguyễn An", "Task 2", ESSAY],
        ["", "", ""],
        ["Trần Bình", "task 1", ESSAY + " Extra."],
    ],
}


async def test_map_sheet_corrects_a_wrong_answer_once():
    llm = FakeLLM(
        [
            {"sheet": "Bài làm", "header_row": 1, "essay_col": 0},  # names, not essays
            {"sheet": "Bài làm", "header_row": 1, "essay_col": 2, "name_col": 0, "task_type_col": 1},
        ]
    )
    mapping = await map_sheet(llm, SHEETS)

    assert mapping.essay_col == 2 and mapping.name_col == 0
    assert "was wrong" in llm.json_calls[1]


async def test_map_sheet_falls_back_to_the_longest_column_without_a_model():
    mapping = await map_sheet(StubExtractor(), SHEETS)

    assert (mapping.sheet, mapping.essay_col, mapping.header_row, mapping.name_col) == ("Bài làm", 2, 1, 0)


def test_apply_mapping_keeps_exact_text_and_row_positions():
    m = SheetMapping(sheet="Bài làm", header_row=1, essay_col=2, name_col=0, task_type_col=1)
    essays, headers = apply_mapping(m, SHEETS)

    assert headers == ["Họ tên", "Đề", "Bài viết"]
    assert [(e.student_label, e.task_type, e.ref["row"]) for e in essays] == [
        ("Nguyễn An", "task_2", 2),
        ("Trần Bình", "task_1", 4),
    ]
    assert essays[1].essay_text == ESSAY + " Extra."


def test_guess_mapping_rejects_an_empty_sheet():
    with pytest.raises(ValueError):
        guess_mapping({"S": [[""]]})


@pytest.mark.parametrize(
    "value,expected",
    [("Task 1", "task_1"), ("task_2", "task_2"), ("2", "task_2"), ("Room 21", None), ("", None)],
)
def test_normalize_task_type(value, expected):
    assert normalize_task_type(value) == expected


# ------------------------------------------------------------------ documents

DOC = "\n".join(["Class 10A", "Student: An", ESSAY, "Student: Binh", ESSAY + " Second."])


async def test_split_document_cuts_exact_lines():
    llm = FakeLLM(
        [{"essays": [{"start_line": 2, "end_line": 2, "student_label": "An"}, {"start_line": 4, "end_line": 4, "student_label": "Binh", "task_type": "task_2"}]}]
    )
    essays = await split_document(llm, DOC, "class.docx")

    assert [e.student_label for e in essays] == ["An", "Binh"]
    assert essays[1].essay_text == ESSAY + " Second." and essays[1].task_type == "task_2"
    assert "L2: " in llm.json_calls[0]


async def test_split_document_keeps_the_whole_text_when_the_model_answers_nonsense():
    llm = FakeLLM([{"essays": [{"start_line": 9, "end_line": 99}]}])
    essays = await split_document(llm, DOC, "class.docx")

    assert len(essays) == 1 and essays[0].essay_text == DOC


# ---------------------------------------------------------------- images, pdf


async def test_photos_of_one_essay_become_one_essay_with_the_visible_question():
    llm = FakeLLM(pages=["Page one text.", "Page two text."])
    essays = await essay_from_images(llm, [(b"a", "image/jpeg"), (b"b", "image/jpeg")], "IMG_01")

    assert essays[0].essay_text == "Page one text.\n\nPage two text."
    assert essays[0].task_prompt == "Q?" and essays[0].student_label == "IMG_01"


async def test_extract_pdf_sends_only_scanned_pages_to_ocr():
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Typed page with a perfectly fine text layer here.")
    doc.new_page()
    pdf = base64.b64encode(doc.tobytes()).decode()

    llm = FakeLLM(answers=[NotImplementedError()], pages=["Handwritten page."])
    out = await extract(llm, {"kind": "pdf", "fileName": "scan.pdf", "files": [{"data": pdf, "mime": "application/pdf"}]})

    assert llm.transcribe_calls == 1
    assert "Typed page" in out["essays"][0]["essayText"] and "Handwritten page." in out["essays"][0]["essayText"]


async def test_worker_fails_fast_on_an_unreadable_file():
    worker = ExtractionWorker(llm=StubExtractor())
    slept = []
    worker.sleep = lambda s: slept.append(s)

    result = await worker.run_with_retries({"fileId": "f1", "kind": "csv", "files": [{"data": "", "mime": "text/csv"}]})

    assert result["status"] == "failed" and "empty" in result["error"]
    assert slept == []
