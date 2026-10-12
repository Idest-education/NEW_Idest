"""Pure file readers. No network, no LLM: bytes in, text or page images out."""

import csv
import io
import os
import subprocess
import tempfile
from dataclasses import dataclass

# A PDF page with less text than this is treated as a scan and sent to OCR.
MIN_TEXT_CHARS_PER_PAGE = 40
PDF_RENDER_DPI = 150


@dataclass
class PdfPage:
    text: str
    image: bytes | None  # PNG, set only for pages without a usable text layer


def read_pdf(data: bytes, max_pages: int = 60) -> list[PdfPage]:
    import pymupdf

    pages = []
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.page_count > max_pages:
            raise ValueError(f"PDF has {doc.page_count} pages; the limit is {max_pages}")
        for page in doc:
            text = page.get_text().strip()
            if len(text) >= MIN_TEXT_CHARS_PER_PAGE:
                pages.append(PdfPage(text=text, image=None))
            else:
                pages.append(PdfPage(text="", image=page.get_pixmap(dpi=PDF_RENDER_DPI).tobytes("png")))
    return pages


def read_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()


def read_doc(data: bytes) -> str:
    """Legacy Word 97-2003 via antiword (installed in the Docker image)."""
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as f:
        f.write(data)
        path = f.name
    try:
        out = subprocess.run(["antiword", "-w", "0", path], capture_output=True, timeout=60, check=True)
        return out.stdout.decode("utf-8", errors="replace").strip()
    except FileNotFoundError as exc:
        raise ValueError("antiword is not installed; cannot read .doc files") from exc
    except subprocess.CalledProcessError as exc:
        raise ValueError(f"Could not read the .doc file: {exc.stderr.decode(errors='replace')[:200]}") from exc
    finally:
        os.unlink(path)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_sheets(data: bytes, kind: str) -> dict[str, list[list[str]]]:
    """Every sheet as rows of strings. `kind` is xls, xlsx or csv."""
    if kind == "csv":
        text = data.decode("utf-8-sig", errors="replace")
        return {"Sheet1": [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]}

    from python_calamine import CalamineWorkbook

    workbook = CalamineWorkbook.from_filelike(io.BytesIO(data))
    sheets = {}
    for name in workbook.sheet_names:
        rows = workbook.get_sheet_by_name(name).to_python(skip_empty_area=False)
        sheets[name] = [[_cell(v) for v in row] for row in rows]
    return sheets
